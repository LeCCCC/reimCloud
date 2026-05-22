from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path.home() / "Desktop" / "项目文档资料"
TEMPLATE = BASE / "API接口接口文档模版.docx"
OUT = BASE / "差旅报销单系统RESTful接口文档.docx"


def clear_doc(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=10.5, bold=False, color=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_p(doc, text="", style=None, size=10.5, bold=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        style_run(r, size=size, bold=bold)
    return p


def set_table_borders(table, color="A6A6A6", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(4.0), Cm(3.3), Cm(2.0), Cm(7.0), Cm(5.2)]
    hdr = table.rows[0].cells
    for i, text in enumerate(["字段名", "类型", "是否必传", "备注", "格式"]):
        hdr[i].text = text
        hdr[i].width = widths[i]
        set_cell_shading(hdr[i], "D9EAF7")
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                style_run(run, size=9, bold=True, color="1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val or "")
            cells[i].width = widths[i]
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if i == 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    style_run(run, size=8.3, bold=False)
        if row[0] and row[0] == row[1] == row[2] == row[3] == row[4]:
            for c in cells:
                set_cell_shading(c, "F2F2F2")
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        style_run(run, size=8.5, bold=True, color="404040")
    set_table_borders(table)
    add_p(doc)
    return table


def group(name):
    return [name, name, name, name, name]


COMMON_RESULT = [
    ["code", "String", "Y", "响应编码，0表示成功，非0表示失败", ""],
    ["message", "String", "Y", "响应消息", ""],
    ["data", "Object", "N", "业务响应数据", ""],
]


def api(doc, title, method_name, scene, path, req, resp, base="ip+端口+/api/v1"):
    add_p(doc, title, "Heading 2", size=13, bold=True)
    add_p(doc, f"接口调用地址：{base}")
    add_p(doc, f"方法名:{method_name}")
    add_p(doc, f"应用场景：{scene}")
    add_p(doc, f"接口路径:{path}")
    add_p(doc, "接口入参：")
    add_table(doc, req)
    add_p(doc, "接口出参：")
    add_table(doc, resp)


def main():
    doc = Document(TEMPLATE)
    clear_doc(doc)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name not in styles:
            continue
        style = styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)

    title = add_p(doc, "差旅报销单系统RESTful接口文档", "Title", size=20, bold=True)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = add_p(doc, "依据《差旅报销单概要设计》编制，接口统一采用 JSON 请求与响应，日期时间采用 yyyy-MM-dd HH:mm:ss。", size=10)
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_p(doc, "1.报销单管理", "Heading 2", size=14, bold=True)

    api(
        doc,
        "1.1查询报销单分页列表",
        "listTravelReimbursements",
        "报销单列表页面根据报销单号、标题、事由、费用归属公司、报销部门、报销人、业务类型、单据状态等条件分页查询报销单数据。",
        "GET /travel-reimbursements",
        [
            ["current", "Integer", "Y", "当前页，默认1", ""],
            ["size", "Integer", "Y", "每页大小，默认10", ""],
            ["billNo", "String", "N", "报销单号，支持模糊查询", ""],
            ["reimbursementTitle", "String", "N", "报销标题，支持模糊查询", "最大500字"],
            ["businessTripReason", "String", "N", "出差事由，支持模糊查询", "最大500字"],
            ["reimCompanyId", "String", "N", "费用归属公司ID", ""],
            ["reimDepartmentId", "String", "N", "报销部门ID", ""],
            ["reimburserId", "String", "N", "报销人ID", ""],
            ["businessTypeId", "String", "N", "业务类型ID", ""],
            ["billStatus", "String", "N", "单据状态", "0草稿 1已完成 2已作废"],
        ],
        [
            ["total", "Integer", "Y", "总条数", ""],
            ["pages", "Integer", "Y", "总页数", ""],
            ["current", "Integer", "Y", "当前页", ""],
            ["size", "Integer", "Y", "每页大小", ""],
            ["records", "List<TravelReimbursementListVO>", "Y", "报销单列表集合", ""],
            group("TravelReimbursementListVO"),
            ["id", "String", "Y", "主键ID", ""],
            ["billNo", "String", "Y", "报销单号", ""],
            ["billStatus", "String", "Y", "单据状态编码", "0草稿 1已完成 2已作废"],
            ["billStatusName", "String", "Y", "单据状态名称", ""],
            ["reimburserNo", "String", "Y", "报销人工号", ""],
            ["reimburserName", "String", "Y", "报销人姓名", ""],
            ["reimDepartmentNo", "String", "Y", "报销部门编号", ""],
            ["reimDepartmentName", "String", "Y", "报销部门名称", ""],
            ["reimCompanyName", "String", "Y", "费用归属公司名称", ""],
            ["businessTypeName", "String", "Y", "业务类型名称", ""],
            ["reimbursementTitle", "String", "Y", "报销标题", ""],
            ["businessTripReason", "String", "Y", "报销事由", ""],
            ["subsidyTotal", "BigDecimal", "Y", "补助金额", "保留2位小数"],
            ["creationTime", "String", "Y", "创建时间", "yyyy-MM-dd HH:mm:ss"],
        ],
    )

    save_req = [
        ["data", "TravelReimbursementSaveDTO", "Y", "报销单保存对象", ""],
        group("TravelReimbursementSaveDTO"),
        ["reimbursementTitle", "String", "Y", "报销标题", "最大500字"],
        ["businessTripReason", "String", "Y", "出差事由", "最大500字"],
        ["reimburserId", "String", "Y", "报销人ID", ""],
        ["reimburserNo", "String", "Y", "报销人工号", ""],
        ["reimburserName", "String", "Y", "报销人姓名", ""],
        ["reimDepartmentId", "String", "Y", "报销部门ID", ""],
        ["reimDepartmentNo", "String", "Y", "报销部门编号", ""],
        ["reimDepartmentName", "String", "Y", "报销部门名称", ""],
        ["reimCompanyId", "String", "Y", "费用归属公司ID", ""],
        ["reimCompanyNo", "String", "Y", "费用归属公司编号", ""],
        ["reimCompanyName", "String", "Y", "费用归属公司名称", ""],
        ["businessTypeId", "String", "Y", "业务类型ID", ""],
        ["businessTypeNo", "String", "Y", "业务类型编号", ""],
        ["businessTypeName", "String", "Y", "业务类型名称", ""],
        ["tripList", "List<TripDTO>", "N", "补录行程集合；草稿可为空，提交必须至少一条", ""],
        ["subsidyList", "List<SubsidyDTO>", "N", "补助信息集合，通常由行程自动生成后保存", ""],
        ["allocationList", "List<AllocationDTO>", "Y", "费用归属及分摊集合", ""],
        ["remarks", "String", "N", "备注信息", "最大1000字"],
    ]

    api(
        doc,
        "1.2新增报销单",
        "createTravelReimbursement",
        "报销单详情页面新增保存草稿，保存基础信息、补录行程、补助信息、费用分摊及备注信息。",
        "POST /travel-reimbursements",
        save_req,
        COMMON_RESULT
        + [
            group("TravelReimbursementCreateResult"),
            ["id", "String", "Y", "新生成报销单主键ID", ""],
            ["billNo", "String", "Y", "新生成报销单号", ""],
            ["billStatus", "String", "Y", "单据状态", "0草稿"],
        ],
    )

    api(
        doc,
        "1.3查询报销单详情",
        "getTravelReimbursementDetail",
        "点击报销单号或报销标题进入单据详情页面，回显基础信息、补录行程、补助日历、费用合计、分摊及备注。",
        "GET /travel-reimbursements/{id}",
        [["id", "String", "Y", "报销单主键ID", "Path"]],
        COMMON_RESULT
        + [
            group("TravelReimbursementDetailVO"),
            ["id", "String", "Y", "报销单主键ID", ""],
            ["billNo", "String", "Y", "报销单号", ""],
            ["billStatus", "String", "Y", "单据状态", "0草稿 1已完成 2已作废"],
            ["creationTime", "String", "Y", "创建时间", "yyyy-MM-dd HH:mm:ss"],
            ["reimbursementTitle", "String", "Y", "报销标题", ""],
            ["businessTripReason", "String", "Y", "出差事由", ""],
            ["reimburserId", "String", "Y", "报销人ID", ""],
            ["reimburserNo", "String", "Y", "报销人工号", ""],
            ["reimburserName", "String", "Y", "报销人姓名", ""],
            ["reimDepartmentId", "String", "Y", "报销部门ID", ""],
            ["reimDepartmentName", "String", "Y", "报销部门名称", ""],
            ["reimCompanyId", "String", "Y", "费用归属公司ID", ""],
            ["reimCompanyName", "String", "Y", "费用归属公司名称", ""],
            ["businessTypeId", "String", "Y", "业务类型ID", ""],
            ["businessTypeName", "String", "Y", "业务类型名称", ""],
            ["tripList", "List<TripVO>", "Y", "补录行程列表", ""],
            ["subsidyList", "List<SubsidyVO>", "Y", "补助信息列表", ""],
            ["costSummary", "CostSummaryVO", "Y", "费用合计信息", ""],
            ["allocationList", "List<AllocationVO>", "Y", "费用归属及分摊列表", ""],
            ["remarks", "String", "N", "备注信息", ""],
        ],
    )

    api(
        doc,
        "1.4修改报销单",
        "updateTravelReimbursement",
        "编辑草稿状态报销单，保存页面当前全部数据，并在后台执行必填、行程重复、金额范围、分摊比例与金额合计校验。",
        "PUT /travel-reimbursements/{id}",
        [["id", "String", "Y", "报销单主键ID", "Path"]] + save_req,
        COMMON_RESULT
        + [
            group("TravelReimbursementUpdateResult"),
            ["id", "String", "Y", "报销单主键ID", ""],
            ["billStatus", "String", "Y", "单据状态", "0草稿"],
            ["updateTime", "String", "Y", "更新时间", "yyyy-MM-dd HH:mm:ss"],
        ],
    )

    api(
        doc,
        "1.5提交报销单",
        "submitTravelReimbursement",
        "点击提交按钮时提交报销单。后台需校验页面必填项、人员+日期范围不可重复、补助金额不大于标准金额、分摊比例合计为100%、分摊金额合计等于补助总金额。",
        "POST /travel-reimbursements/{id}/submission",
        [["id", "String", "Y", "报销单主键ID", "Path"]],
        COMMON_RESULT
        + [
            group("SubmitResult"),
            ["id", "String", "Y", "报销单主键ID", ""],
            ["billNo", "String", "Y", "报销单号", ""],
            ["billStatus", "String", "Y", "提交后单据状态", "1已完成"],
            ["submitTime", "String", "Y", "提交时间", "yyyy-MM-dd HH:mm:ss"],
        ],
    )

    api(
        doc,
        "1.6作废报销单",
        "voidTravelReimbursement",
        "列表或详情页面对未完成或需废弃的报销单执行作废操作，作废后单据状态为已作废。",
        "DELETE /travel-reimbursements/{id}",
        [["id", "String", "Y", "报销单主键ID", "Path"]],
        COMMON_RESULT
        + [
            group("VoidResult"),
            ["id", "String", "Y", "报销单主键ID", ""],
            ["billStatus", "String", "Y", "作废后单据状态", "2已作废"],
        ],
    )

    add_p(doc, "2.补录行程管理", "Heading 2", size=14, bold=True)

    trip_req = [
        ["travelerId", "String", "Y", "出行人ID", ""],
        ["travelerNo", "String", "Y", "出行人工号", ""],
        ["travelerName", "String", "Y", "出行人姓名", ""],
        ["departureCityNo", "String", "Y", "出发城市编号", ""],
        ["departureCityName", "String", "Y", "出发城市名称", ""],
        ["arrivalCityNo", "String", "Y", "到达城市编号", ""],
        ["arrivalCityName", "String", "Y", "到达城市名称", ""],
        ["departureDate", "String", "Y", "出发日期", "yyyy-MM-dd"],
        ["arrivalDate", "String", "Y", "到达日期，不可早于出发日期且不可晚于当前日期", "yyyy-MM-dd"],
        ["tripDescription", "String", "Y", "行程说明", "最大500字"],
    ]
    trip_resp = COMMON_RESULT + [
        group("TripVO"),
        ["tripId", "String", "Y", "补录行程ID", ""],
        ["travelerId", "String", "Y", "出行人ID", ""],
        ["travelerName", "String", "Y", "出行人姓名", ""],
        ["tripDateRange", "String", "Y", "出差日期范围", "yyyy-MM-dd至yyyy-MM-dd"],
        ["tripRoute", "String", "Y", "行程", "出发城市-到达城市"],
        ["tripDescription", "String", "Y", "行程说明", ""],
        ["subsidyId", "String", "Y", "关联生成的补助信息ID", ""],
    ]
    api(
        doc,
        "2.1新增补录行程",
        "createTrip",
        "补录未从申请单带入或未产生费用的行程信息。保存时按人员+日期范围校验不可重复，并关联生成补助信息及补助日历。",
        "POST /travel-reimbursements/{id}/trips",
        [["id", "String", "Y", "报销单主键ID", "Path"]] + trip_req,
        trip_resp,
    )

    api(
        doc,
        "2.2修改补录行程",
        "updateTrip",
        "编辑补录行程弹窗回显的数据，保存后重新校验人员+日期范围不可重复，并同步刷新关联补助信息和补助日历。",
        "PUT /travel-reimbursements/{id}/trips/{tripId}",
        [["id", "String", "Y", "报销单主键ID", "Path"], ["tripId", "String", "Y", "补录行程ID", "Path"]] + trip_req,
        trip_resp,
    )

    api(
        doc,
        "2.3删除补录行程",
        "deleteTrip",
        "确认删除补录行程，同时删除或作废其关联生成的补助信息及补助日历数据。",
        "DELETE /travel-reimbursements/{id}/trips/{tripId}",
        [["id", "String", "Y", "报销单主键ID", "Path"], ["tripId", "String", "Y", "补录行程ID", "Path"]],
        COMMON_RESULT
        + [
            group("DeleteTripResult"),
            ["id", "String", "Y", "报销单主键ID", ""],
            ["tripId", "String", "Y", "已删除补录行程ID", ""],
            ["subsidyTotal", "BigDecimal", "Y", "删除后补助总金额", "保留2位小数"],
        ],
    )

    add_p(doc, "3.补助信息管理", "Heading 2", size=14, bold=True)

    api(
        doc,
        "3.1查询补助日历",
        "getSubsidyCalendar",
        "打开补助信息编辑弹窗时，根据补录行程的日期范围、到达城市和城市类型查询补助日历。",
        "GET /travel-reimbursements/{id}/subsidies/{subsidyId}/calendar",
        [["id", "String", "Y", "报销单主键ID", "Path"], ["subsidyId", "String", "Y", "补助信息ID", "Path"]],
        COMMON_RESULT
        + [
            group("SubsidyCalendarVO"),
            ["subsidyId", "String", "Y", "补助信息ID", ""],
            ["businessTypeName", "String", "Y", "出差类型，与主单业务类型一致", ""],
            ["startDate", "String", "Y", "开始日期", "yyyy-MM-dd"],
            ["endDate", "String", "Y", "结束日期", "yyyy-MM-dd"],
            ["days", "Integer", "Y", "天数", ""],
            ["route", "String", "Y", "行程地点", "出发城市-到达城市"],
            ["subsidyAmount", "BigDecimal", "Y", "已选补助金额合计", "保留2位小数"],
            ["standardAmount", "BigDecimal", "Y", "已选标准金额合计", "保留2位小数"],
            ["calendarList", "List<SubsidyCalendarItemVO>", "Y", "补助日历明细", ""],
            group("SubsidyCalendarItemVO"),
            ["calendarId", "String", "Y", "补助日历明细ID", ""],
            ["travelDate", "String", "Y", "出差日期", "yyyy-MM-dd"],
            ["weekName", "String", "Y", "星期", ""],
            ["subsidyCityName", "String", "Y", "补助城市", ""],
            ["cityType", "String", "Y", "城市类型", "1一线 2二线 3三线"],
            ["mealSelected", "Boolean", "Y", "餐费补助是否选中", "true/false"],
            ["mealStandard", "BigDecimal", "Y", "餐费补助标准金额", "一线100 二线80 三线50"],
            ["mealAmount", "BigDecimal", "Y", "餐费补助申请金额", "不得大于标准金额"],
            ["transportSelected", "Boolean", "Y", "交通补助是否选中", "true/false"],
            ["transportStandard", "BigDecimal", "Y", "交通补助标准金额", "40元/天"],
            ["transportAmount", "BigDecimal", "Y", "交通补助申请金额", "不得大于标准金额"],
            ["phoneSelected", "Boolean", "Y", "通讯补助是否选中", "true/false"],
            ["phoneStandard", "BigDecimal", "Y", "通讯补助标准金额", "40元/天"],
            ["phoneAmount", "BigDecimal", "Y", "通讯补助申请金额", "不得大于标准金额"],
        ],
    )

    api(
        doc,
        "3.2保存补助日历",
        "saveSubsidyCalendar",
        "保存补助日历中横向日期选择、纵向补助项选择及各补助申请金额，后台校验未选中项金额不可保存、申请金额为正数且不大于标准金额。",
        "PUT /travel-reimbursements/{id}/subsidies/{subsidyId}/calendar",
        [
            ["id", "String", "Y", "报销单主键ID", "Path"],
            ["subsidyId", "String", "Y", "补助信息ID", "Path"],
            ["calendarList", "List<SubsidyCalendarSaveDTO>", "Y", "补助日历保存明细", ""],
            group("SubsidyCalendarSaveDTO"),
            ["calendarId", "String", "Y", "补助日历明细ID", ""],
            ["mealSelected", "Boolean", "Y", "餐费补助是否选中", "true/false"],
            ["mealAmount", "BigDecimal", "Y", "餐费补助申请金额", "保留2位小数"],
            ["transportSelected", "Boolean", "Y", "交通补助是否选中", "true/false"],
            ["transportAmount", "BigDecimal", "Y", "交通补助申请金额", "保留2位小数"],
            ["phoneSelected", "Boolean", "Y", "通讯补助是否选中", "true/false"],
            ["phoneAmount", "BigDecimal", "Y", "通讯补助申请金额", "保留2位小数"],
        ],
        COMMON_RESULT
        + [
            group("SubsidySaveResult"),
            ["subsidyId", "String", "Y", "补助信息ID", ""],
            ["mealAllowance", "BigDecimal", "Y", "餐费补助合计", "保留2位小数"],
            ["transportationAllowance", "BigDecimal", "Y", "交通补助合计", "保留2位小数"],
            ["phoneAllowance", "BigDecimal", "Y", "通讯补助合计", "保留2位小数"],
            ["subsidyTotal", "BigDecimal", "Y", "补助总金额", "保留2位小数"],
        ],
    )

    add_p(doc, "4.费用归属及分摊管理", "Heading 2", size=14, bold=True)

    alloc_req = [
        ["allocationList", "List<AllocationDTO>", "Y", "分摊明细集合，至少保留一条", ""],
        group("AllocationDTO"),
        ["allocationId", "String", "N", "分摊明细ID，新增为空", ""],
        ["reimCompanyId", "String", "Y", "费用归属公司ID", ""],
        ["reimCompanyNo", "String", "Y", "费用归属公司编号", ""],
        ["reimCompanyName", "String", "Y", "费用归属公司名称", ""],
        ["projectId", "String", "Y", "项目ID", ""],
        ["projectNo", "String", "Y", "项目编号", ""],
        ["projectName", "String", "Y", "项目名称", ""],
        ["allocationRatio", "BigDecimal", "Y", "分摊比例，存值区间0-1", "页面展示为百分比，保留2位小数"],
        ["allocationAmount", "BigDecimal", "Y", "分摊金额", "保留2位小数"],
    ]
    api(
        doc,
        "4.1保存费用分摊",
        "saveAllocations",
        "保存费用归属及分摊分区数据，校验至少一条分摊信息、分摊比例合计100%、分摊金额合计等于补助总金额。",
        "PUT /travel-reimbursements/{id}/allocations",
        [["id", "String", "Y", "报销单主键ID", "Path"]] + alloc_req,
        COMMON_RESULT
        + [
            group("AllocationSaveResult"),
            ["subsidyTotal", "BigDecimal", "Y", "补助总金额", "保留2位小数"],
            ["allocationTotalAmount", "BigDecimal", "Y", "分摊金额合计", "保留2位小数"],
            ["allocationTotalRatio", "BigDecimal", "Y", "分摊比例合计", "1.00表示100%"],
            ["allocationList", "List<AllocationVO>", "Y", "保存后的分摊明细", ""],
        ],
    )

    api(
        doc,
        "4.2费用分摊均摊试算",
        "calculateEqualAllocations",
        "点击均摊按钮时，根据分摊数据条数对比例和金额进行均摊，除不尽的差值放在首行。",
        "POST /travel-reimbursements/{id}/allocation-equalizations",
        [
            ["id", "String", "Y", "报销单主键ID", "Path"],
            ["subsidyTotal", "BigDecimal", "Y", "补助总金额", "保留2位小数"],
            ["allocationCount", "Integer", "Y", "参与均摊的分摊行数", "必须大于0"],
        ],
        COMMON_RESULT
        + [
            group("AllocationEqualizationResult"),
            ["allocationList", "List<AllocationVO>", "Y", "均摊后的分摊明细", ""],
            ["allocationRatio", "BigDecimal", "Y", "分摊比例", "区间0-1"],
            ["allocationAmount", "BigDecimal", "Y", "分摊金额", "保留2位小数"],
        ],
    )

    add_p(doc, "5.基础数据管理", "Heading 2", size=14, bold=True)

    base_resp = COMMON_RESULT + [
        group("BaseDataVO"),
        ["id", "String", "Y", "主键ID", ""],
        ["no", "String", "Y", "编号", ""],
        ["name", "String", "Y", "名称", ""],
    ]
    api(doc, "5.1查询费用归属公司列表", "listReimCompanies", "费用归属公司下拉控件数据查询。", "GET /reim-companies", [["keyword", "String", "N", "编号或名称关键字", "Query"]], base_resp)
    api(doc, "5.2查询报销部门列表", "listReimDepartments", "报销部门下拉控件数据查询。", "GET /reim-departments", [["keyword", "String", "N", "编号或名称关键字", "Query"]], base_resp)
    api(doc, "5.3查询员工列表", "listEmployees", "报销人、出行人下拉控件数据查询。", "GET /employees", [["keyword", "String", "N", "工号或姓名关键字", "Query"]], COMMON_RESULT + [group("EmployeeVO"), ["reimburserId", "String", "Y", "员工ID", ""], ["reimburserNo", "String", "Y", "员工工号", ""], ["reimburserName", "String", "Y", "员工姓名", ""]])
    api(doc, "5.4查询业务类型树", "listBusinessTypeTree", "业务类型树形下拉控件数据查询。", "GET /business-types/tree", [["superiorId", "String", "N", "上级业务类型ID，不传返回全量树", "none表示最上级"]], COMMON_RESULT + [group("BusinessTypeVO"), ["businessTypeId", "String", "Y", "业务类型ID", ""], ["businessTypeNo", "String", "Y", "业务类型编号", ""], ["businessTypeName", "String", "Y", "业务类型名称", ""], ["thereSubordinateNode", "String", "Y", "是否有下级节点", "0无 1有"], ["superiorId", "String", "Y", "上级业务类型ID", "none表示最上级"], ["children", "List<BusinessTypeVO>", "N", "下级业务类型集合", ""]])
    api(doc, "5.5查询城市列表", "listCities", "出发城市、到达城市及补助城市控件数据查询。", "GET /cities", [["keyword", "String", "N", "城市编号或名称关键字", "Query"]], COMMON_RESULT + [group("CityVO"), ["cityNo", "String", "Y", "城市编号", ""], ["cityName", "String", "Y", "城市名称", ""], ["cityType", "String", "Y", "城市类型", "1一线 2二线 3三线"]])
    api(doc, "5.6查询项目列表", "listProjects", "费用分摊项目下拉控件数据查询。", "GET /projects", [["keyword", "String", "N", "项目编号或名称关键字", "Query"]], COMMON_RESULT + [group("ProjectVO"), ["projectId", "String", "Y", "项目ID", ""], ["projectNo", "String", "Y", "项目编号", ""], ["projectName", "String", "Y", "项目名称", ""]])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
